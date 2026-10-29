"""Relay 生成的文件（Excel、PPT）：按结构化内容生成，生成后重新打开逐项核对，再交给用户。

“执行成功不等于做对了”：文件生成后用独立的解析器重新读出来，和要写的内容逐格、逐页对一遍，
结果（ok / 不一致的地方）和文件一起给用户看。核对不通过的文件照样给出，但会写明哪里不对。
"""

import io
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_ROWS = 5000
MAX_SLIDES = 40
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


class OutputError(ValueError):
    """内容不合要求（可以原样告诉模型）。"""


def safe_name(name: str, ext: str) -> str:
    base = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", (name or "").strip()).strip(". ") or "Relay"
    base = base[:80]
    return base if base.lower().endswith(f".{ext}") else f"{base}.{ext}"


def _norm(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:19]
    return str(value).strip()


# ---------- Excel ----------


def build_xlsx(sheets: list[dict]) -> bytes:
    """sheets: [{"name", "columns": [...], "rows": [[...]]}]。以“=”开头的字符串写成公式。"""
    if not sheets:
        raise OutputError("至少要有一个工作表")
    wb = Workbook()
    wb.remove(wb.active)
    names = set()
    for spec in sheets:
        name = re.sub(r"[\[\]:*?/\\]", "_", str(spec.get("name") or "Sheet"))[:31] or "Sheet"
        if name in names:
            raise OutputError(f"工作表名重复：{name}")
        names.add(name)
        columns = [str(c) for c in spec.get("columns") or []]
        rows = spec.get("rows") or []
        if len(rows) > MAX_ROWS:
            raise OutputError(f"《{name}》超过 {MAX_ROWS} 行")
        ws = wb.create_sheet(name)
        if columns:
            ws.append(columns)
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="F0EEE6")
                cell.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
        for row in rows:
            if columns and len(row) > len(columns):
                raise OutputError(f"《{name}》有一行比表头多出 {len(row) - len(columns)} 列")
            ws.append(list(row))
        for index in range(1, max(len(columns), max((len(r) for r in rows), default=0)) + 1):
            width = max([len(_norm(columns[index - 1])) if index <= len(columns) else 0]
                        + [len(_norm(r[index - 1])) for r in rows[:200] if index <= len(r)])
            ws.column_dimensions[get_column_letter(index)].width = min(max(10, width * 1.6 + 2), 60)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def check_xlsx(data: bytes, sheets: list[dict]) -> dict:
    """重新打开文件，逐个工作表、逐格对照要写的内容。"""
    wb = load_workbook(io.BytesIO(data))  # 不取计算结果：公式按原文比较
    problems, cells = [], 0
    for spec in sheets:
        name = re.sub(r"[\[\]:*?/\\]", "_", str(spec.get("name") or "Sheet"))[:31] or "Sheet"
        if name not in wb.sheetnames:
            problems.append(f"缺少工作表《{name}》")
            continue
        ws = wb[name]
        expected = ([list(spec["columns"])] if spec.get("columns") else []) + [list(r) for r in spec.get("rows") or []]
        if ws.max_row != max(len(expected), 1) and expected:
            problems.append(f"《{name}》应有 {len(expected)} 行，实际 {ws.max_row} 行")
        for r, row in enumerate(expected, start=1):
            for c, want in enumerate(row, start=1):
                cells += 1
                have = ws.cell(row=r, column=c).value
                if _norm(have) != _norm(want):
                    problems.append(f"《{name}》{get_column_letter(c)}{r} 应为“{_norm(want)[:20]}”，实际“{_norm(have)[:20]}”")
    extra = [n for n in wb.sheetnames if n not in {re.sub(r'[\[\]:*?/\\]', '_', str(s.get('name') or 'Sheet'))[:31] for s in sheets}]
    if extra:
        problems.append("多出工作表：" + "、".join(extra))
    summary = "；".join(f"《{s.get('name')}》{len(s.get('rows') or [])} 行" for s in sheets)
    return {"ok": not problems, "detail": (f"重新打开文件核对了 {len(sheets)} 个工作表、{cells} 个单元格，与要求一致（{summary}）"
                                            if not problems else "核对不一致：" + "；".join(problems[:6]))}


# ---------- PPT ----------


def build_pptx(slides: list[dict], title: str | None = None) -> bytes:
    """slides: [{"title", "bullets": [...], "notes", "table": {"columns": [...], "rows": [[...]]}}]。第一页可用 title 做封面。"""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    if not slides:
        raise OutputError("至少要有一页")
    if len(slides) > MAX_SLIDES:
        raise OutputError(f"超过 {MAX_SLIDES} 页")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    if title:
        cover = prs.slides.add_slide(prs.slide_layouts[0])
        cover.shapes.title.text = title
    for spec in slides:
        table = spec.get("table")
        layout = prs.slide_layouts[5] if table else prs.slide_layouts[1]  # 只有标题 / 标题和正文
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = str(spec.get("title") or "")
        bullets = [str(b) for b in spec.get("bullets") or []]
        if table:
            columns = [str(c) for c in table.get("columns") or []]
            rows = table.get("rows") or []
            if not columns:
                raise OutputError(f"《{spec.get('title')}》的表格没有表头")
            shape = slide.shapes.add_table(len(rows) + 1, len(columns), Inches(0.6), Inches(1.6), Inches(12.1), Inches(0.4) * (len(rows) + 1))
            for c, name in enumerate(columns):
                shape.table.cell(0, c).text = name
            for r, row in enumerate(rows, start=1):
                for c in range(len(columns)):
                    shape.table.cell(r, c).text = _norm(row[c]) if c < len(row) else ""
            if bullets:
                box = slide.shapes.add_textbox(Inches(0.6), Inches(1.6) + Inches(0.4) * (len(rows) + 2), Inches(12.1), Inches(1))
                box.text_frame.text = bullets[0]
                for b in bullets[1:]:
                    box.text_frame.add_paragraph().text = b
        else:
            body = slide.placeholders[1].text_frame
            body.text = bullets[0] if bullets else ""
            for b in bullets[1:]:
                body.add_paragraph().text = b
            for p in body.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(20)
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(spec["notes"])
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def check_pptx(data: bytes, slides: list[dict], title: str | None = None) -> dict:
    """重新打开文件，逐页对照标题、要点和表格。"""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    actual = list(prs.slides)
    offset = 1 if title else 0
    problems = []
    if len(actual) != len(slides) + offset:
        problems.append(f"应有 {len(slides) + offset} 页，实际 {len(actual)} 页")
    if title and actual and (actual[0].shapes.title.text if actual[0].shapes.title else "") != title:
        problems.append("封面标题不一致")
    for n, spec in enumerate(slides):
        if n + offset >= len(actual):
            break
        slide = actual[n + offset]
        page = n + offset + 1
        if (slide.shapes.title.text if slide.shapes.title else "") != str(spec.get("title") or ""):
            problems.append(f"第 {page} 页标题不一致")
        text = "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
        missing = [b for b in spec.get("bullets") or [] if str(b) not in text]
        if missing:
            problems.append(f"第 {page} 页缺少要点：{str(missing[0])[:20]}")
        if spec.get("table"):
            tables = [s.table for s in slide.shapes if s.has_table]
            if not tables:
                problems.append(f"第 {page} 页缺少表格")
            else:
                t = tables[0]
                want = [list(spec["table"]["columns"])] + [list(r) for r in spec["table"].get("rows") or []]
                for r, row in enumerate(want):
                    for c, value in enumerate(row):
                        if r >= len(t.rows) or c >= len(t.columns) or t.cell(r, c).text != _norm(value):
                            problems.append(f"第 {page} 页表格第 {r + 1} 行第 {c + 1} 列不一致")
                            break
    return {"ok": not problems, "detail": (f"重新打开文件核对了 {len(actual)} 页的标题、要点和表格，与要求一致"
                                            if not problems else "核对不一致：" + "；".join(problems[:6]))}


# ---------- Word ----------

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_SECTIONS = 60


def build_docx(title: str | None, sections: list[dict]) -> bytes:
    """sections: [{"heading", "paragraphs": [...], "bullets": [...], "table": {"columns", "rows"}}]。"""
    from docx import Document

    if not sections:
        raise OutputError("至少要有一节内容")
    if len(sections) > MAX_SECTIONS:
        raise OutputError(f"超过 {MAX_SECTIONS} 节")
    doc = Document()
    if title:
        doc.add_heading(title, level=0)
    for spec in sections:
        if spec.get("heading"):
            doc.add_heading(str(spec["heading"]), level=1)
        for text in spec.get("paragraphs") or []:
            doc.add_paragraph(str(text))
        for text in spec.get("bullets") or []:
            doc.add_paragraph(str(text), style="List Bullet")
        table = spec.get("table")
        if table:
            columns = [str(c) for c in table.get("columns") or []]
            if not columns:
                raise OutputError(f"《{spec.get('heading')}》的表格没有表头")
            rows = table.get("rows") or []
            grid = doc.add_table(rows=len(rows) + 1, cols=len(columns))
            grid.style = "Table Grid"
            for c, name in enumerate(columns):
                grid.cell(0, c).text = name
            for r, row in enumerate(rows, start=1):
                if len(row) > len(columns):
                    raise OutputError(f"《{spec.get('heading')}》的表格有一行比表头多出 {len(row) - len(columns)} 列")
                for c in range(len(columns)):
                    grid.cell(r, c).text = _norm(row[c]) if c < len(row) else ""
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def check_docx(data: bytes, title: str | None, sections: list[dict]) -> dict:
    """重新打开文件：标题、每节的标题和段落、要点、表格逐项对照（顺序也要对）。"""
    from docx import Document

    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    expected = ([title] if title else [])
    for spec in sections:
        expected += ([str(spec["heading"])] if spec.get("heading") else []) + [str(x) for x in spec.get("paragraphs") or []] \
            + [str(x) for x in spec.get("bullets") or []]
    problems, position = [], 0
    for want in expected:
        try:
            position = texts.index(want, position) + 1
        except ValueError:
            problems.append(f"缺少或顺序不对：“{want[:20]}”")
    tables = [s["table"] for s in sections if s.get("table")]
    if len(doc.tables) != len(tables):
        problems.append(f"应有 {len(tables)} 张表，实际 {len(doc.tables)} 张")
    for n, (want, have) in enumerate(zip(tables, doc.tables, strict=False), start=1):
        rows = [list(want["columns"])] + [list(r) for r in want.get("rows") or []]
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                if r >= len(have.rows) or c >= len(have.columns) or have.cell(r, c).text != _norm(value):
                    problems.append(f"第 {n} 张表第 {r + 1} 行第 {c + 1} 列不一致")
                    break
    return {"ok": not problems, "detail": (f"重新打开文件核对了 {len(expected)} 段文字和 {len(tables)} 张表，与要求一致"
                                            if not problems else "核对不一致：" + "；".join(problems[:6]))}


MEDIA = {"xlsx": XLSX, "pptx": PPTX, "docx": DOCX}
FEISHU_FILE_TYPE = {"xlsx": "xls", "pptx": "ppt", "docx": "doc"}


def build(kind: str, spec: dict) -> tuple[bytes, dict]:
    """按种类生成并核对，返回 (文件, 核对结果)。"""
    if kind == "xlsx":
        data = build_xlsx(spec["sheets"])
        return data, check_xlsx(data, spec["sheets"])
    if kind == "pptx":
        data = build_pptx(spec["slides"], spec.get("title"))
        return data, check_pptx(data, spec["slides"], spec.get("title"))
    if kind == "docx":
        data = build_docx(spec.get("title"), spec["sections"])
        return data, check_docx(data, spec.get("title"), spec["sections"])
    raise OutputError(f"不支持的文件类型：{kind}")


def size_of(kind: str, spec: dict) -> int:
    """粗略的工作量：决定当场生成，还是交给后台做完再告诉他。"""
    if kind == "xlsx":
        return sum(len(s.get("rows") or []) * max(1, len(s.get("columns") or [])) for s in spec["sheets"])
    if kind == "pptx":
        return len(spec["slides"]) * 200
    return sum(len(s.get("paragraphs") or []) + len(s.get("bullets") or []) + len((s.get("table") or {}).get("rows") or []) * 5
               for s in spec["sections"]) * 20
