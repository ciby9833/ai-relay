"""把文件原件（邮件附件、云空间上传的文件）转成模型能读的文字。

支持 xlsx / docx / pdf / csv / txt。表格保留行号和列字母（如 S49），
Word 保留修订（[+新增+] [-删除-]）和批注，方便对照“谁改了什么、谁提了什么意见”。
"""

import csv
import io
import re
import zipfile
from collections import OrderedDict
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

MAX_BYTES = 60 * 1024 * 1024
PAGE_CHARS = 60_000
CELL_CHARS = 1500

HEX_NOISE = re.compile(r"\b[0-9A-F]{64,}\b")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class FileError(Exception):
    pass


def kind_of(filename: str, content_type: str | None = None) -> str | None:
    name = (filename or "").lower()
    for ext in ("xlsx", "xlsm", "docx", "pdf", "csv", "txt", "md"):
        if name.endswith("." + ext):
            return {"xlsm": "xlsx", "md": "txt"}.get(ext, ext)
    ct = content_type or ""
    if "spreadsheetml" in ct:
        return "xlsx"
    if "wordprocessingml" in ct:
        return "docx"
    if "pdf" in ct:
        return "pdf"
    if ct.startswith("text/"):
        return "txt"
    return None


def _cell(value) -> str:
    text = str(value).strip() if value is not None else ""
    text = re.sub(r"\s*\n\s*", " ⏎ ", text)
    return text if len(text) <= CELL_CHARS else text[:CELL_CHARS] + "…"


def xlsx_text(data: bytes) -> str:
    """每个工作表一段；每行写成 `R49 | A: … | S: …`，空格子跳过。合并单元格注明范围。"""
    try:
        book = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001 - 损坏或加密的文件
        raise FileError(f"无法解析这个 Excel 文件：{exc}") from exc
    parts = []
    for sheet in book.worksheets:
        merged = {str(r).split(":")[0]: str(r) for r in sheet.merged_cells.ranges}
        rows = []
        for row in sheet.iter_rows():
            cells = []
            for c in row:
                text = _cell(c.value)
                if text:
                    ref = f"{get_column_letter(c.column)}{c.row}"
                    span = f"（合并 {merged[ref]}）" if ref in merged else ""
                    cells.append(f"{get_column_letter(c.column)}{span}: {text}")
            if cells:
                rows.append(f"R{row[0].row} | " + " | ".join(cells))
        state = "（隐藏）" if sheet.sheet_state != "visible" else ""
        parts.append(f"## 工作表「{sheet.title}」{state}，共 {sheet.max_row} 行 {sheet.max_column} 列\n" + ("\n".join(rows) or "（空）"))
    return "\n\n".join(parts)


def _runs_text(node) -> str:
    """一组 run 的文字；跳过隐藏文字（w:vanish，常见于文档水印/追踪码）。"""
    out = []
    runs = [node] if node.tag == W + "r" else list(node.iter(W + "r"))
    for run in runs:
        props = run.find(W + "rPr")
        if props is not None and props.find(W + "vanish") is not None:
            continue
        for el in run:
            if el.tag in (W + "t", W + "delText") and el.text:
                out.append(el.text)
            elif el.tag == W + "tab":
                out.append("\t")
            elif el.tag in (W + "br", W + "cr"):
                out.append("\n")
    return HEX_NOISE.sub("", "".join(out))


def _paragraph(p) -> str:
    """一段文字，修订标出来：插入 [+…+]，删除 [-…-]。"""
    out = []
    for child in p:
        if child.tag == W + "ins":
            text = _runs_text(child)
            out.append(f"[+{text}+]" if text else "")
        elif child.tag == W + "del":
            text = _runs_text(child)
            out.append(f"[-{text}-]" if text else "")
        elif child.tag == W + "commentRangeStart":
            out.append(f"〔批注{child.get(W + 'id')}→")
        elif child.tag == W + "commentRangeEnd":
            out.append("〕")
        elif child.tag in (W + "r", W + "hyperlink", W + "smartTag"):
            out.append(_runs_text(child))
    return "".join(out).strip()


def docx_text(data: bytes) -> str:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        body = ElementTree.fromstring(archive.read("word/document.xml")).find(W + "body")
    except Exception as exc:  # noqa: BLE001
        raise FileError(f"无法解析这个 Word 文件：{exc}") from exc
    lines = []
    for block in body:
        if block.tag == W + "p":
            text = _paragraph(block)
            if text:
                lines.append(text)
        elif block.tag == W + "tbl":
            for tr in block.iter(W + "tr"):
                cells = [" ".join(filter(None, (_paragraph(p) for p in tc.iter(W + "p")))) for tc in tr.iter(W + "tc")]
                if any(cells):
                    lines.append("| " + " | ".join(cells) + " |")
    text = "\n".join(lines)
    if "word/comments.xml" in archive.namelist():
        comments = []
        for c in ElementTree.fromstring(archive.read("word/comments.xml")).iter(W + "comment"):
            body_text = " ".join(filter(None, (_paragraph(p) for p in c.iter(W + "p"))))
            comments.append(f"- 批注{c.get(W + 'id')}（{c.get(W + 'author') or '?'}）：{body_text}")
        if comments:
            text += "\n\n## 批注\n" + "\n".join(comments)
    if "[+" in text or "[-" in text:
        text = "（文档含修订：[+新增+] [-删除-]；〔批注N→…〕标出批注对应的原文）\n\n" + text
    return text


def pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [f"## 第 {i} 页\n{(page.extract_text() or '').strip()}" for i, page in enumerate(reader.pages, 1)]
    except Exception as exc:  # noqa: BLE001
        raise FileError(f"无法解析这个 PDF：{exc}") from exc
    text = "\n\n".join(pages)
    if len(re.sub(r"\s|## 第 \d+ 页", "", text)) < 20:
        raise FileError("这个 PDF 没有可提取的文字（可能是扫描件），暂时读不了")
    return text


def csv_text(data: bytes) -> str:
    raw = data.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(raw)))
    return "\n".join(
        f"R{i} | " + " | ".join(f"{get_column_letter(j)}: {_cell(v)}" for j, v in enumerate(row, 1) if _cell(v))
        for i, row in enumerate(rows, 1) if any(v.strip() for v in row)
    )


def to_text(data: bytes, filename: str, content_type: str | None = None) -> str:
    if len(data) > MAX_BYTES:
        raise FileError(f"文件太大（{len(data) // 1024 // 1024}MB），超过 {MAX_BYTES // 1024 // 1024}MB 的读取上限")
    kind = kind_of(filename, content_type)
    if kind == "xlsx":
        return xlsx_text(data)
    if kind == "docx":
        return docx_text(data)
    if kind == "pdf":
        return pdf_text(data)
    if kind == "csv":
        return csv_text(data)
    if kind == "txt":
        return data.decode("utf-8", errors="replace")
    raise FileError(f"暂不支持读取这种文件：{filename}（支持 xlsx、docx、pdf、csv、txt）")


def page(text: str, offset: int = 0, limit: int = PAGE_CHARS) -> str:
    """长文件分段给模型：明确告诉它读到哪了、怎么继续。"""
    chunk = text[offset:offset + limit]
    end = offset + len(chunk)
    if offset == 0 and end >= len(text):
        return chunk
    note = f"[全文 {len(text)} 字，这里是第 {offset}–{end} 字。"
    note += f"继续读传 offset={end}。]" if end < len(text) else "已读到结尾。]"
    return f"{note}\n\n{chunk}"


class Cache:
    """解析过的文件留在内存里（同一份大 Excel 翻页时不重复下载解析）。"""

    def __init__(self, size: int = 8):
        self.size, self.items = size, OrderedDict()

    def get(self, key):
        if key in self.items:
            self.items.move_to_end(key)
            return self.items[key]
        return None

    def put(self, key, value):
        self.items[key] = value
        self.items.move_to_end(key)
        while len(self.items) > self.size:
            self.items.popitem(last=False)


cache = Cache()
