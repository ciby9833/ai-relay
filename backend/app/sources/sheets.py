"""飞书电子表格：按单元格读写。

读：值按纯文本取（ToString），整理成 `R行号 | 列字母: 值`，模型说“第 49 行”时有据可查。
写：只按单元格填值（例如别人发来的反馈表），准备时先检查编辑权限、合并单元格和原值，确认时再核对原值。
"""

import re
from urllib.parse import parse_qs, urlparse

from openpyxl.utils import column_index_from_string, get_column_letter

from app.feishu.client import FeishuClient, FeishuError
from app.models import User

CELL = re.compile(r"^([A-Z]{1,3})([1-9][0-9]{0,5})$")
RANGE = re.compile(r"^([A-Z]{1,3})([1-9][0-9]{0,5}):([A-Z]{1,3})([1-9][0-9]{0,5})$")
TOKEN = re.compile(r"^[A-Za-z0-9_-]{5,128}$")
CELL_CHARS = 2000
MAX_CELLS = 5000
UNSUPPORTED = "#UNSUPPORT VALUE"


class SheetError(Exception):
    pass


def parse_target(spreadsheet: str, sheet_id: str | None = None) -> tuple[str, str | None]:
    """表格链接或 token → (token, sheet_id)；链接里的 ?sheet= 作为默认工作表。"""
    spreadsheet = spreadsheet.strip()
    if spreadsheet.startswith("http"):
        parsed = urlparse(spreadsheet)
        match = re.fullmatch(r"/sheets/([A-Za-z0-9_-]+)/?", parsed.path)
        if not match:
            raise SheetError("这不是飞书电子表格链接（/sheets/…）；知识库里的表格先用 read_feishu_doc 取到 spreadsheet_token")
        sheet_id = sheet_id or (parse_qs(parsed.query).get("sheet") or [None])[0]
        spreadsheet = match.group(1)
    if not TOKEN.match(spreadsheet):
        raise SheetError("spreadsheet_token 格式不对")
    if sheet_id is not None and not TOKEN.match(sheet_id) and not re.fullmatch(r"[A-Za-z0-9]{1,32}", sheet_id):
        raise SheetError("sheet_id 格式不对")
    return spreadsheet, sheet_id


def parse_cell(ref: str) -> tuple[int, int]:
    """'S49' → (行 49, 列 19)。"""
    match = CELL.match(ref.strip().upper())
    if not match:
        raise SheetError(f"单元格写法不对：{ref}（应为列字母+行号，例如 S49）")
    return int(match.group(2)), column_index_from_string(match.group(1))


def cell_name(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, list):  # 富文本片段
        value = "".join(seg.get("text", "") for seg in value if isinstance(seg, dict))
    elif isinstance(value, dict):
        value = value.get("text") or value.get("link") or ""
    text = str(value)
    return "[图片或不支持的内容]" if text == UNSUPPORTED else text


def explain(exc: FeishuError, *, writing: bool = False) -> str:
    """把飞书的报错翻译成用户能处理的说法，分清“应用没开权限”和“你对这个文件没有权限”。"""
    message = str(exc)
    if exc.code == 99991679 or "required privileges" in message or "scope" in message.lower():
        return ("Relay 应用还没有开通电子表格的" + ("编辑" if writing else "读取") + "权限，需要管理员在飞书后台开通"
                "（sheets:spreadsheet）并发布，然后重新登录 Relay")
    if exc.code in (1310213, 1310217, 91403, 1061004, 403) or "permission" in message.lower() or "forbidden" in message.lower():
        return "你对这个表格没有" + ("编辑" if writing else "查看") + "权限，需要找表格所有者开通后再试"
    if exc.code in (1310214, 1310211, 404) or "not found" in message.lower():
        return "找不到这个表格或工作表，请核对链接和 sheet_id"
    return message


async def sheets_of(feishu: FeishuClient, user: User, token: str) -> list[dict]:
    data = await feishu.user_api(user, "GET", f"/sheets/v3/spreadsheets/{token}/sheets/query")
    return data.get("sheets") or []


async def sheet_info(feishu: FeishuClient, user: User, token: str, sheet_id: str | None) -> dict:
    sheets = await sheets_of(feishu, user, token)
    if not sheets:
        raise SheetError("这个表格里没有工作表")
    if sheet_id is None:
        visible = [s for s in sheets if not s.get("hidden")]
        if len(visible) > 1:
            names = "、".join(f"{s['title']}（{s['sheet_id']}）" for s in visible)
            raise SheetError(f"表格有多个工作表，请指定 sheet_id：{names}")
        return (visible or sheets)[0]
    for sheet in sheets:
        if sheet["sheet_id"] == sheet_id:
            return sheet
    raise SheetError(f"没有找到工作表 {sheet_id}")


def merges(sheet: dict) -> list[tuple[int, int, int, int]]:
    """合并区域 → (起始行, 起始列, 结束行, 结束列)，1 起、含结束。"""
    return [
        (m["start_row_index"] + 1, m["start_column_index"] + 1, m["end_row_index"] + 1, m["end_column_index"] + 1)
        for m in sheet.get("merges") or []
    ]


def merge_of(row: int, col: int, areas: list[tuple[int, int, int, int]]) -> tuple[int, int, int, int] | None:
    for area in areas:
        r0, c0, r1, c1 = area
        if r0 <= row <= r1 and c0 <= col <= c1 and (r0, c0) != (r1, c1):
            return area
    return None


async def read_range(feishu: FeishuClient, user: User, token: str, sheet_id: str, ref: str) -> str:
    """读一个范围，返回 `R行号 | 列: 值` 的文字表格。"""
    match = RANGE.match(ref.strip().upper())
    if not match:
        raise SheetError("范围写法不对，应为 A1:T50 这样的形式")
    c0, r0 = column_index_from_string(match.group(1)), int(match.group(2))
    c1, r1 = column_index_from_string(match.group(3)), int(match.group(4))
    if (r1 - r0 + 1) * (c1 - c0 + 1) > MAX_CELLS:
        raise SheetError(f"一次最多读 {MAX_CELLS} 格，请缩小范围")
    data = await feishu.user_api(user, "GET", f"/sheets/v2/spreadsheets/{token}/values/{sheet_id}!{ref.upper()}",
                                 params={"valueRenderOption": "ToString"})
    values = (data.get("valueRange") or {}).get("values") or []
    lines = []
    for i, row in enumerate(values):
        cells = []
        for j, value in enumerate(row):
            text = _text(value).strip()
            if text:
                text = re.sub(r"\s*\n\s*", " ⏎ ", text)
                cells.append(f"{get_column_letter(c0 + j)}: {text[:CELL_CHARS]}{'…' if len(text) > CELL_CHARS else ''}")
        if cells:
            lines.append(f"R{r0 + i} | " + " | ".join(cells))
    return "\n".join(lines) or "（这个范围是空的）"


async def read_cells(feishu: FeishuClient, user: User, token: str, sheet_id: str, cells: list[str]) -> dict[str, str]:
    """一组单元格的当前值（纯文本）。"""
    result: dict[str, str] = {}
    for start in range(0, len(cells), 100):
        chunk = cells[start:start + 100]
        data = await feishu.user_api(
            user, "GET", f"/sheets/v2/spreadsheets/{token}/values_batch_get",
            params={"ranges": ",".join(f"{sheet_id}!{c}:{c}" for c in chunk), "valueRenderOption": "ToString"},
        )
        for cell, value_range in zip(chunk, data.get("valueRanges") or [], strict=False):
            values = value_range.get("values") or [[None]]
            result[cell] = _text(values[0][0] if values and values[0] else None)
    return result


async def can_edit(feishu: FeishuClient, user: User, token: str) -> bool | None:
    """用户本人对这个表格有没有编辑权限；查不了时返回 None（写的时候再看）。"""
    try:
        data = await feishu.user_api(user, "GET", f"/drive/v1/permissions/{token}/members/auth",
                                     params={"type": "sheet", "action": "edit"})
        return bool(data.get("auth_result"))
    except FeishuError:
        return None


async def prepare_cells(feishu: FeishuClient, user: User, token: str, sheet_id: str | None, cells: list[dict]) -> dict:
    """校验并整理要填的单元格：合并单元格换成左上角，读出原值。返回待确认操作的 payload。"""
    if not cells:
        raise SheetError("没有要填写的单元格")
    if len(cells) > 200:
        raise SheetError("一次最多填 200 个单元格")
    editable = await can_edit(feishu, user, token)
    if editable is False:
        raise SheetError("你对这个表格只有查看权限，不能写入。需要找表格所有者开编辑权限；也可以把要填的内容整理给你自己粘贴")
    sheet = await sheet_info(feishu, user, token, sheet_id)
    areas = merges(sheet)
    notes, targets = [], {}
    for entry in cells:
        row, col = parse_cell(entry["cell"])
        value = entry["value"]
        if isinstance(value, str) and len(value) > 20000:
            raise SheetError(f"{entry['cell']} 的内容太长")
        if area := merge_of(row, col, areas):
            anchor = cell_name(area[0], area[1])
            span = f"{anchor}:{cell_name(area[2], area[3])}"
            if (row, col) != (area[0], area[1]):
                notes.append(f"{entry['cell']} 在合并单元格 {span} 里，改为写入左上角 {anchor}")
            else:
                notes.append(f"{anchor} 是合并单元格 {span}，内容显示在整个合并区域")
            row, col = area[0], area[1]
        name = cell_name(row, col)
        if name in targets and targets[name] != value:
            raise SheetError(f"{name} 被指定了两个不同的值")
        targets[name] = value
    before = await read_cells(feishu, user, token, sheet["sheet_id"], list(targets))
    return {
        "operation": "fill_sheet_cells", "spreadsheet_token": token, "sheet_id": sheet["sheet_id"],
        "sheet_title": sheet.get("title"),
        "cells": [{"cell": c, "before": before.get(c, ""), "after": v} for c, v in targets.items()],
        "notes": notes, "editable_checked": editable is True,
    }


def preview(payload: dict) -> str:
    lines = [f"在工作表「{payload.get('sheet_title') or payload['sheet_id']}」填写 {len(payload['cells'])} 个单元格："]
    for c in payload["cells"]:
        before = c["before"].strip() or "（空）"
        lines.append(f"\n{c['cell']}\n  原值：{before[:300]}{'…' if len(before) > 300 else ''}\n  新值：{c['after']}")
    if payload.get("notes"):
        lines.append("\n注意：" + "；".join(payload["notes"]))
    return "\n".join(lines)


async def write_cells(feishu: FeishuClient, user: User, payload: dict) -> None:
    """确认后写入：先核对原值没被别人改过，再一次性写入。"""
    token, sheet_id = payload["spreadsheet_token"], payload["sheet_id"]
    cells = payload["cells"]
    current = await read_cells(feishu, user, token, sheet_id, [c["cell"] for c in cells])
    changed = [c["cell"] for c in cells if current.get(c["cell"], "") != c["before"]]
    if changed:
        raise SheetError(f"{'、'.join(changed)} 在你确认前已被修改，本次没有写入。请重新读取后再准备")
    await feishu.user_api(
        user, "POST", f"/sheets/v2/spreadsheets/{token}/values_batch_update",
        json={"valueRanges": [{"range": f"{sheet_id}!{c['cell']}:{c['cell']}", "values": [[c["after"]]]} for c in cells]},
    )
