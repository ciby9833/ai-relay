"""生成 Excel / PPT：生成后重新打开逐项核对；被改过的文件核对得出来；只有本人能下载；飞书里对话时把文件发给他。"""
import io
import os
import uuid

import pytest
from openpyxl import load_workbook

from app import outputs

db = pytest.mark.skipif(not os.getenv("RUN_DB_TESTS"), reason="需要 RUN_DB_TESTS=1 和数据库")

SHEETS = [{"name": "网点延误", "columns": ["网点", "延误单数", "占比"], "rows": [["Medan", 12, "40%"], ["Batam", 18, "60%"], ["合计", "=SUM(B2:B3)", ""]]}]
SLIDES = [
    {"title": "本周延误概况", "bullets": ["Medan 延误 12 单", "Batam 延误 18 单"], "notes": "先讲结论"},
    {"title": "明细", "table": {"columns": ["网点", "延误单数"], "rows": [["Medan", 12], ["Batam", 18]]}},
]


def test_excel_is_built_and_checked_cell_by_cell():
    data = outputs.build_xlsx(SHEETS)
    result = outputs.check_xlsx(data, SHEETS)
    assert result["ok"] and "12 个单元格" in result["detail"]
    wb = load_workbook(io.BytesIO(data))
    assert wb["网点延误"]["B4"].value == "=SUM(B2:B3)" and wb["网点延误"].freeze_panes == "A2"
    # 文件被改过（或生成出了问题）：核对得出来，指出哪一格
    wb["网点延误"]["B2"] = 21
    buffer = io.BytesIO()
    wb.save(buffer)
    bad = outputs.check_xlsx(buffer.getvalue(), SHEETS)
    assert not bad["ok"] and "B2 应为“12”，实际“21”" in bad["detail"]
    with pytest.raises(outputs.OutputError):
        outputs.build_xlsx([{"name": "x", "columns": ["a"], "rows": [[1, 2]]}])


def test_ppt_is_built_and_checked_page_by_page():
    from pptx import Presentation

    data = outputs.build_pptx(SLIDES, "延误复盘")
    result = outputs.check_pptx(data, SLIDES, "延误复盘")
    assert result["ok"] and "3 页" in result["detail"]
    prs = Presentation(io.BytesIO(data))
    prs.slides[1].shapes.title.text = "别的标题"
    buffer = io.BytesIO()
    prs.save(buffer)
    bad = outputs.check_pptx(buffer.getvalue(), SLIDES, "延误复盘")
    assert not bad["ok"] and "第 2 页标题不一致" in bad["detail"]
    assert outputs.safe_name("周报/延误:复盘", "pptx") == "周报_延误_复盘.pptx"


@db
async def test_generated_files_reach_only_the_owner():
    import httpx
    from sqlalchemy import delete

    from app import deps
    from app.agent.tools.base import ToolContext, execute
    from app.db import SessionLocal, engine
    from app.main import app
    from app.models import Conversation, User

    class Feishu:
        def __init__(self):
            self.files = []

        async def send_file(self, open_id, data, filename, file_type):
            self.files.append((open_id, filename, file_type, len(data)))
            return "om_file"

    tag = uuid.uuid4().hex[:8]
    async with SessionLocal() as session:
        me = User(feishu_open_id=f"ou_{tag}", name="Xavier(吴川)", timezone="Asia/Jakarta")
        other = User(feishu_open_id=f"ou_o_{tag}", name="Kange", timezone="Asia/Jakarta")
        session.add_all([me, other])
        await session.flush()
        ids = [me.id, other.id]
        conversation = Conversation(user_id=me.id, channel="feishu")
        session.add(conversation)
        await session.flush()
        feishu = Feishu()
        ctx = ToolContext(session=session, user=me, conversation=conversation, feishu=feishu)
        try:
            text, failed = await execute(ctx, "create_excel_file", {"filename": "延误/明细", "sheets": SHEETS})
            assert not failed and "核对通过" in text and feishu.files == [(me.feishu_open_id, "延误_明细.xlsx", "xls", feishu.files[0][3])]
            card = ctx.cards[-1]
            assert card["kind"] == "file" and card["verified"]["ok"]
            text, failed = await execute(ctx, "create_ppt_file", {"filename": "延误复盘", "title": "延误复盘", "slides": SLIDES})
            assert not failed and "核对通过" in text and feishu.files[-1][2] == "ppt"
            await session.commit()
            users = {u.id: u for u in [await session.get(User, i) for i in ids]}
            app.dependency_overrides[deps.get_current_user] = lambda: users[ids[0]]
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://relay") as client:
                response = await client.get(f"/api/files/{card['file_id']}")
                assert response.status_code == 200 and response.headers["content-type"] == outputs.XLSX
                assert outputs.check_xlsx(response.content, SHEETS)["ok"]
                app.dependency_overrides[deps.get_current_user] = lambda: users[ids[1]]
                assert (await client.get(f"/api/files/{card['file_id']}")).status_code == 404  # 别人下载不了
        finally:
            app.dependency_overrides.pop(deps.get_current_user, None)
            await session.rollback()
            await session.execute(delete(User).where(User.id.in_(ids)))
            await session.commit()
            await engine.dispose()
