"""文件解析、按单元格填表、卡片说明的位置、飞书只发最后一段。不访问真实飞书。"""
import io
import zipfile
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.agent.loop import CARD_NOTE, _with_card_notes, strip_leaked_notes
from app.feishu.client import FeishuError
from app.sources import files, sheets

USER = SimpleNamespace(id="test", feishu_open_id="ou_me")


def test_xlsx_keeps_row_numbers_columns_and_merges():
    book = Workbook()
    ws = book.active
    ws.title = "反馈"
    ws["A1"], ws["S1"] = "条款", "Status 状态"
    ws["A49"], ws["T49"] = "4. Data Protection", "同意\n第二行"
    ws.merge_cells("S154:S160")
    ws["S154"] = "Pending"
    buf = io.BytesIO()
    book.save(buf)
    text = files.to_text(buf.getvalue(), "V9 vs v10.xlsx")
    assert "R49 | A: 4. Data Protection | T: 同意 ⏎ 第二行" in text
    assert "S（合并 S154:S160）: Pending" in text


def test_docx_marks_tracked_changes_comments_and_skips_hidden_text():
    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    document = f"""<w:document {w}><w:body>
      <w:p><w:r><w:t>Old </w:t></w:r><w:del><w:r><w:delText>term</w:delText></w:r></w:del><w:ins><w:r><w:t>clause</w:t></w:r></w:ins>
        <w:r><w:rPr><w:vanish/></w:rPr><w:t>SECRET</w:t></w:r></w:p>
      <w:p><w:commentRangeStart w:id="0"/><w:r><w:t>BCE system</w:t></w:r><w:commentRangeEnd w:id="0"/></w:p>
    </w:body></w:document>"""
    comments = f'<w:comments {w}><w:comment w:id="0" w:author="LEI"><w:p><w:r><w:t>Same as LEL</w:t></w:r></w:p></w:comment></w:comments>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", document)
        z.writestr("word/comments.xml", comments)
    text = files.to_text(buf.getvalue(), "contract.docx")
    assert "Old [-term-][+clause+]" in text and "SECRET" not in text
    assert "〔批注0→BCE system〕" in text and "批注0（LEI）：Same as LEL" in text


def test_unsupported_file_and_paging():
    with pytest.raises(files.FileError, match="暂不支持"):
        files.to_text(b"x", "logo.png")
    text = "a" * 150
    first = files.page(text, 0, 100)
    assert "offset=100" in first and files.page(text, 100, 100).endswith("a" * 50)
    assert files.page("short") == "short"


class SheetFeishu:
    def __init__(self, before, editable=True, fail=None):
        self.before, self.editable, self.fail = before, editable, fail
        self.writes = []

    async def user_api(self, user, method, path, **kwargs):
        if path.endswith("/members/auth"):
            return {"auth_result": self.editable}
        if path.endswith("/sheets/query"):
            return {"sheets": [{"sheet_id": "s1", "title": "Sheet1", "merges": [
                {"start_row_index": 153, "end_row_index": 159, "start_column_index": 18, "end_column_index": 18}]}]}
        if path.endswith("/values_batch_get"):
            cells = [r.split("!")[1].split(":")[0] for r in kwargs["params"]["ranges"].split(",")]
            return {"valueRanges": [{"values": [[self.before.get(c)]]} for c in cells]}
        if path.endswith("/values_batch_update"):
            if self.fail:
                raise self.fail
            self.writes.append(kwargs["json"])
            return {}
        raise AssertionError(path)


async def test_fill_cells_checks_permission_merges_and_stale_values():
    target = sheets.parse_target("https://x.sg.feishu.cn/sheets/Y8S4s1v2ohC5EH?from=copy&sheet=s1")
    assert target == ("Y8S4s1v2ohC5EH", "s1")

    with pytest.raises(sheets.SheetError, match="只有查看权限"):
        await sheets.prepare_cells(SheetFeishu({}, editable=False), USER, *target, [{"cell": "S49", "value": "Agreed"}])

    f = SheetFeishu({"S154": "Pending Discussion"})
    payload = await sheets.prepare_cells(f, USER, *target, [{"cell": "S49", "value": "Agreed"}, {"cell": "S157", "value": "Pending"}])
    assert [c["cell"] for c in payload["cells"]] == ["S49", "S154"]
    assert payload["cells"][1]["before"] == "Pending Discussion"
    assert "S157 在合并单元格 S154:S160 里" in payload["notes"][0]
    assert "原值：（空）" in sheets.preview(payload) and f.writes == []

    f.before["S49"] = "他人改过"
    with pytest.raises(sheets.SheetError, match="已被修改"):
        await sheets.write_cells(f, USER, payload)
    assert f.writes == []

    f.before["S49"] = None
    await sheets.write_cells(f, USER, payload)
    assert f.writes[0]["valueRanges"][0] == {"range": "s1!S49:S49", "values": [["Agreed"]]}


def test_errors_distinguish_app_scope_from_file_permission():
    assert "管理员" in sheets.explain(FeishuError("required privileges [sheets:spreadsheet]", 99991679), writing=True)
    assert "表格所有者" in sheets.explain(FeishuError("forbidden", 1310213), writing=True)


def test_card_notes_go_to_next_user_message_not_assistant():
    rows = [
        SimpleNamespace(role="user", content=[{"type": "text", "text": "找一下"}]),
        SimpleNamespace(role="assistant", content=[
            {"type": "tool_use", "id": "t1", "name": "x", "input": {}},
            {"type": "card", "card": {"kind": "doc", "title": "合同"}},
        ]),
        SimpleNamespace(role="user", content=[{"type": "tool_result", "tool_use_id": "t1", "content": "ok"}]),
        SimpleNamespace(role="assistant", content=[{"type": "text", "text": "好了"}, {"type": "card", "card": {"kind": "doc", "title": "周报"}}]),
    ]
    messages = _with_card_notes(rows, {}, {})
    assert all(b["type"] != "card" for m in messages for b in m["content"])
    assert not any(CARD_NOTE in b.get("text", "") for m in messages if m["role"] == "assistant" for b in m["content"])
    assert messages[2]["content"][0]["type"] == "tool_result" and "合同" in messages[2]["content"][1]["text"]
    assert messages[-1]["role"] == "user" and "周报" in messages[-1]["content"][0]["text"]


def test_leaked_card_note_is_stripped():
    leaked = "（以下内容以卡片形式展示给了用户）\n[doc] 合同 片段\n[doc] 合同 片段2\n结论：可以同意"
    assert strip_leaked_notes(leaked) == "结论：可以同意"


def test_interrupted_tool_turn_is_closed_before_new_message():
    """上一轮在工具调用后被打断：新消息不能并进工具结果那一条（Gemini 会拒绝），要先补一句助手的话。"""
    from types import SimpleNamespace

    from app.agent.loop import _with_card_notes, sanitize_history

    rows = [
        SimpleNamespace(role="user", content=[{"type": "text", "text": "查一下"}]),
        SimpleNamespace(role="assistant", content=[{"type": "tool_use", "id": "t1", "name": "x", "input": {}}]),
        SimpleNamespace(role="user", content=[{"type": "tool_result", "tool_use_id": "t1", "name": "x", "content": "ok"}]),
    ]
    history = sanitize_history(_with_card_notes(rows, {}, {}))
    assert history[-1]["role"] == "user" and history[-1]["content"][0]["type"] == "tool_result"
